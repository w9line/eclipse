import re
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional, Iterator, Dict, Any

from PIL import Image
from PIL.ExifTags import TAGS
import pypdf
import docx
import openpyxl
from pptx import Presentation

warnings.filterwarnings("ignore", category=UserWarning, module="docx")
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


@dataclass
class Finding:
    source: str
    path: str
    kind: str
    excerpt: str
    start: int = 0
    end: int = 0
    entropy: Optional[float] = None
    category: str = "metadata"
    severity: str = "low"
    hint: Optional[str] = None


def _is_binary_data(content: bytes, sample_size: int = 1024) -> bool:
    sample = content[:sample_size]
    return b"\x00" in sample


def _scan_text_for_metadata(text: str, path: str) -> Iterator[Finding]:
    for match in re.finditer(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text):
        yield Finding(
            source="workdir",
            path=path,
            kind="email_in_text",
            excerpt=match.group(),
            severity="low",
            hint="Обнаружен email в исходном коде или логах. Убедитесь, что это не персональные данные."
        )

    internal_patterns = [
        r"\b[A-Za-z0-9.-]*\.(local|corp|intranet|internal)\b",
        r"\b(dev|staging|test|qa)[.-][A-Za-z0-9.-]+\b",
        r"\b192\.168\.\d{1,3}\.\d{1,3}\b",
        r"\b10\.\d{1,3}\.\d{1,3}\.\d{1,3}\b",
        r"\b172\.(1[6-9]|2[0-9]|3[01])\.\d{1,3}\.\d{1,3}\b",
    ]
    for pattern in internal_patterns:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            yield Finding(
                source="workdir",
                path=path,
                kind="internal_network_artifact",
                excerpt=match.group(),
                severity="medium",
                hint="Найдена ссылка на внутреннюю сеть. Это может помочь злоумышленнику в рекогносцировке."
            )

    user_path_re = r"[\\/](home|Users|user|users)[\\/][A-Za-z0-9_-]{3,}"
    for match in re.finditer(user_path_re, text):
        username = match.group().split("/")[-1].split("\\")[-1]
        yield Finding(
            source="workdir",
            path=path,
            kind="username_in_path",
            excerpt=f"User path: {match.group()}",
            severity="low",
            hint=f"Обнаружено имя пользователя ОС: {username}. Может быть использовано для атак."
        )

    debug_comments = [
        r"//\s*TODO.*",
        r"//\s*FIXME.*",
        r"<!--.*debug.*-->",
        r"#\s*DEBUG.*",
        r"console\.log\(",
        r"print\(",
        r"log\(",
    ]
    for pattern in debug_comments:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            yield Finding(
                source="workdir",
                path=path,
                kind="debug_artifact",
                excerpt=match.group()[:100],
                severity="info",
                hint="Найден отладочный след в коде. В продакшене это нежелательно."
            )


def _scan_docx(path: Path) -> Iterator[Finding]:
    try:
        doc = docx.Document(path)
        core_props = doc.core_properties
        meta = {}
        if core_props.author:
            meta["author"] = core_props.author
        if core_props.company:
            meta["company"] = core_props.company
        if core_props.comments:
            meta["comments"] = core_props.comments
        if core_props.category:
            meta["category"] = core_props.category
        if core_props.last_modified_by:
            meta["last_modified_by"] = core_props.last_modified_by

        for key, value in meta.items():
            if value:
                yield Finding(
                    source="workdir",
                    path=str(path),
                    kind=f"docx_{key}",
                    excerpt=f"{key}: {value}",
                    severity="low",
                    hint="Метаданные документа Word могут раскрывать внутреннюю информацию."
                )
    except Exception:
        pass 


def _scan_xlsx(path: Path) -> Iterator[Finding]:
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        props = wb.properties
        meta = {}
        if props.creator:
            meta["creator"] = props.creator
        if props.lastModifiedBy:
            meta["last_modified_by"] = props.lastModifiedBy
        if props.title:
            meta["title"] = props.title
        if props.description:
            meta["description"] = props.description
        if props.subject:
            meta["subject"] = props.subject

        for key, value in meta.items():
            if value:
                yield Finding(
                    source="workdir",
                    path=str(path),
                    kind=f"xlsx_{key}",
                    excerpt=f"{key}: {value}",
                    severity="low",
                    hint="Метаданные Excel-файла могут содержать служебную информацию."
                )
    except Exception:
        pass


def _scan_pptx(path: Path) -> Iterator[Finding]:
    try:
        prs = Presentation(path)
        core_props = prs.core_properties
        meta = {}
        if core_props.author:
            meta["author"] = core_props.author
        if core_props.company:
            meta["company"] = core_props.company
        if core_props.comments:
            meta["comments"] = core_props.comments

        for key, value in meta.items():
            if value:
                yield Finding(
                    source="workdir",
                    path=str(path),
                    kind=f"pptx_{key}",
                    excerpt=f"{key}: {value}",
                    severity="low",
                    hint="Метаданные презентации могут раскрывать внутренние данные."
                )
    except Exception:
        pass


def _scan_pdf(path: Path) -> Iterator[Finding]:
    try:
        reader = pypdf.PdfReader(path)
        if not reader.metadata:
            return
        meta = reader.metadata
        for key, value in meta.items():
            if value and isinstance(value, str) and value.strip():
                clean_key = key.lstrip("/")
                yield Finding(
                    source="workdir",
                    path=str(path),
                    kind=f"pdf_{clean_key.lower()}",
                    excerpt=f"{clean_key}: {value}",
                    severity="low",
                    hint="PDF-метаданные могут содержать автора, организацию, ПО создания."
                )
    except Exception:
        pass


def _scan_image_exif(path: Path) -> Iterator[Finding]:
    try:
        with Image.open(path) as img:
            exifdata = img.getexif()
            if not exifdata:
                return
            for tag_id, value in exifdata.items():
                tag = TAGS.get(tag_id, tag_id)
                if not value or not isinstance(value, (str, int, tuple)):
                    continue
                if tag == "GPSInfo":
                    yield Finding(
                        source="workdir",
                        path=str(path),
                        kind="exif_gps",
                        excerpt="GPS coordinates embedded",
                        severity="medium",
                        hint="В изображении есть геолокация! Удалите метаданные перед публикацией."
                    )
                elif tag in ("Artist", "Author", "UserComment", "Copyright", "Software", "Make", "Model"):
                    yield Finding(
                        source="workdir",
                        path=str(path),
                        kind=f"exif_{tag.lower()}",
                        excerpt=f"{tag}: {value}",
                        severity="low",
                        hint="EXIF-метаданные могут раскрывать устройство, ПО или автора."
                    )
    except Exception:
        pass


META_EXTENSIONS = {
    ".docx": _scan_docx,
    ".xlsx": _scan_xlsx,
    ".pptx": _scan_pptx,
    ".pdf": _scan_pdf,
    ".jpg": _scan_image_exif,
    ".jpeg": _scan_image_exif,
    ".png": _scan_image_exif,
}


def scan_file_for_metadata(
    file_path: Path,
    source: str = "workdir",
    max_size: int = 15_000_000  ) -> List[Finding]:
 
    findings: List[Finding] = []

    try:
        if file_path.stat().st_size > max_size:
            return findings

        suffix = file_path.suffix.lower()
        if suffix in META_EXTENSIONS:
            handler = META_EXTENSIONS[suffix]
            findings.extend(handler(file_path))
            try:
                text = file_path.read_text(encoding="utf-8", errors="ignore")
                findings.extend(_scan_text_for_metadata(text, str(file_path)))
            except Exception:
                pass

        else:
            try:
                content = file_path.read_bytes()
                if _is_binary_data(content):
                    return findings
                text = content.decode(errors="ignore")
                findings.extend(_scan_text_for_metadata(text, str(file_path)))
            except Exception:
                pass

    except Exception:
        pass

    for f in findings:
        f.source = source

    return findings